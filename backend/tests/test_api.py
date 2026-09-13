"""API-boundary tests for the FastAPI layer.

Written in pytest style, per the instruction to use pytest in the
actual repository. These require `fastapi`, `httpx`, and `pytest` to
be installed — none of which are available in the sandbox this backend
was built in (no network access to install them). They are written
and reviewed carefully but could not be executed here; run them with
`pytest backend/tests/test_api.py -v` once dependencies are installed
locally. See the final report for the full explanation.

These tests intentionally do NOT re-verify the allocation policy
itself (test_engine.py already does that thoroughly) — they verify
that the API boundary (status codes, response shape, camelCase keys,
error handling) behaves as specified on top of an engine we already
trust.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)


def test_get_plan_returns_healthy_baseline():
    response = client.get("/api/plan")
    assert response.status_code == 200
    body = response.json()
    assert body["dataHealth"] == "healthy"


def test_get_plan_baseline_kpi_values():
    body = client.get("/api/plan").json()
    kpis = body["kpis"]

    assert kpis["expectedPlanT"] == 600.0
    assert kpis["actualReceivedT"] == 560.0
    assert kpis["stationCapacityT"] == 500
    assert kpis["actualBySegment"] == {"A": 90, "B": 160, "C": 180, "D": 130}
    assert kpis["exportT"] == 500
    assert kpis["exportRate"] == pytest.approx(0.893, abs=0.0005)
    assert kpis["localVolumeT"] == 60
    assert kpis["exportRevenueEur"] == 549_500
    assert kpis["localValueEur"] == pytest.approx(4_500)
    assert kpis["totalValueEur"] == pytest.approx(554_000)
    assert kpis["atRiskCount"] == 3


def test_get_plan_baseline_client_outcomes():
    body = client.get("/api/plan").json()
    by_id = {c["clientId"]: c for c in body["clientResults"]}

    assert by_id["C02"]["status"] == "PARTIAL"
    assert by_id["C02"]["reason"] == "INSUFFICIENT_COMPATIBLE_SEGMENT"
    assert by_id["C09"]["status"] == "PARTIAL"
    assert by_id["C09"]["reason"] == "INSUFFICIENT_COMPATIBLE_SEGMENT"
    assert by_id["C08"]["status"] == "PARTIAL"
    assert by_id["C08"]["reason"] == "STATION_CAPACITY_REACHED"


def test_get_plan_invalid_workbook_returns_422(monkeypatch, tmp_path):
    """Point the route at a deliberately broken workbook (a farm whose
    expected mix doesn't sum to 1.0) and confirm the 422 shape."""
    import openpyxl

    from app.core.loader import (
        CLIENTS_HEADER,
        FARMS_HEADER,
        SEGMENT_PRICE_HEADER,
        STATION_HEADER,
    )

    wb = openpyxl.Workbook()
    farms_ws = wb.active
    farms_ws.title = "Farms"
    farms_ws.append(["Farms sheet", None])
    farms_ws.append(["description"])
    farms_ws.append([])
    farms_ws.append(list(FARMS_HEADER))
    # expected mix sums to 0.9, not 1.0 — deliberately invalid
    farms_ws.append(["F01", "Farm One", 20.0, 0.5, 0.4, 0.0, 0.0, 10, 5, 0, 0])

    clients_ws = wb.create_sheet("Clients")
    clients_ws.append(list(CLIENTS_HEADER))
    clients_ws.append(["C01", "Client One", "EXACT", "A", 5, 1000])

    station_ws = wb.create_sheet("Station")
    station_ws.append(list(STATION_HEADER))
    station_ws.append(["STATION-01", 100, 0.1])
    station_ws.append([])
    station_ws.append(list(SEGMENT_PRICE_HEADER))
    for seg, price in (("A", 1000), ("B", 800), ("C", 600), ("D", 400)):
        station_ws.append([seg, price])

    broken_path = tmp_path / "broken.xlsx"
    wb.save(broken_path)

    monkeypatch.setattr("app.api.plan.DATA_PATH", Path(broken_path))

    response = client.get("/api/plan")
    assert response.status_code == 422
    body = response.json()
    assert body["dataHealth"] == "invalid"
    assert any(err["sheet"] == "Farms" and err["id"] == "F01" for err in body["validationErrors"])


def test_assistant_at_risk_clients_returns_grounded_citations():
    response = client.post("/api/assistant", json={"question": "at_risk_clients"})
    assert response.status_code == 200
    body = response.json()

    assert body["available"] is True
    assert set(body["citedIds"]) == {"C02", "C09", "C08"}
    assert body["source"] == "fallback"


def test_assistant_unsupported_question_is_honest():
    response = client.post("/api/assistant", json={"freeText": "Can you book a delivery truck for tomorrow?"})
    assert response.status_code == 200
    body = response.json()

    assert body["available"] is False
    assert body["reason"] == "not_supported_by_current_plan"
    # Must not fabricate an answer under the "unavailable" shape.
    assert "citedIds" not in body or body.get("citedIds") in (None, [])


def test_assistant_rejects_arbitrary_text_in_question_field():
    """`question` is a programmatic chip selector (Literal-typed), not
    a place for typed text — arbitrary strings there must be rejected
    at the request-validation boundary (422), not silently fall through
    to "unsupported"."""
    response = client.post("/api/assistant", json={"question": "what is the risk"})
    assert response.status_code == 422


def test_assistant_rejects_both_question_and_free_text():
    """Sending both fields at once is an ambiguous request and must be
    rejected outright rather than the route silently picking one."""
    response = client.post(
        "/api/assistant",
        json={"question": "at_risk_clients", "freeText": "also tell me about farms"},
    )
    assert response.status_code == 422


def test_assistant_rejects_neither_field_present():
    response = client.post("/api/assistant", json={})
    assert response.status_code == 422