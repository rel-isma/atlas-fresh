"""Workbook loader.

This module's only job is mechanical extraction: turn worksheet cells
into raw Python dicts. It contains zero business rules — it does not
check ranges, does not check sums, does not reject anything. All
judgment happens in validation.py, which consumes this module's output.

The workbook is opened read_only=True: this module is structurally
incapable of writing back to the source file.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl

REQUIRED_SHEETS = ("Farms", "Clients", "Station")

FARMS_HEADER = (
    "farm_id",
    "farm_name",
    "expected_daily_capacity_t",
    "expected_A_pct",
    "expected_B_pct",
    "expected_C_pct",
    "expected_D_pct",
    "actual_A_t",
    "actual_B_t",
    "actual_C_t",
    "actual_D_t",
)
CLIENTS_HEADER = (
    "client_id",
    "client_name",
    "acceptance_mode",
    "requested_segment",
    "demand_t",
    "export_price_per_t_eur",
)
STATION_HEADER = (
    "station_id",
    "export_conditioning_capacity_t",
    "local_market_ratio",
)
SEGMENT_PRICE_HEADER = ("segment", "reference_export_price_per_t_eur")

# How many rows we scan looking for a header. The segment price table on
# the Station sheet sits well below the station row, so this must be
# generous rather than assuming a fixed offset.
HEADER_SCAN_WINDOW = 25
RawRow = dict[str, object]


@dataclass(frozen=True)
class LoaderIssue:
    """A structural problem found while loading (missing sheet/header).

    Kept separate from validation.ValidationIssue because these are not
    business-rule violations — the file itself doesn't have the shape
    we expect.
    """

    sheet: str
    message: str


@dataclass(frozen=True)
class RawWorkbook:
    farms: list[RawRow]
    clients: list[RawRow]
    station: RawRow | None
    segment_prices: dict[str, object]  # segment letter -> raw price value
    loader_issues: list[LoaderIssue]


def _find_header_row(
    rows: list[tuple[object, ...]], header: tuple[str, ...]
) -> int | None:
    """Return the 0-based index of the row matching `header` on its
    leading cells, or None if not found within the scan window."""
    n = len(header)
    for i, row in enumerate(rows[:HEADER_SCAN_WINDOW]):
        if tuple(row[:n]) == header:
            return i
    return None


def _read_rows_after_header(
    rows: list[tuple[object, ...]],
    header_idx: int,
    header: tuple[str, ...],
) -> list[RawRow]:
    """Read every row after the header until the first fully-blank row."""
    out: list[RawRow] = []
    for row in rows[header_idx + 1 :]:
        values = [row[i] if i < len(row) else None for i in range(len(header))]
        if all(v is None for v in values):
            break
        out.append(dict(zip(header, values)))
    return out


def load_workbook(path: str | Path) -> RawWorkbook:
    """Open the workbook read-only and extract raw rows for each sheet.

    Never mutates or writes the source file. Never interprets a cell's
    business meaning — that is validation.py's job.
    """
    issues: list[LoaderIssue] = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for name in REQUIRED_SHEETS:
            if name not in wb.sheetnames:
                issues.append(LoaderIssue(sheet=name, message=f"Sheet '{name}' is missing from the workbook"))

        farms: list[RawRow] = []
        clients: list[RawRow] = []
        station: RawRow | None = None
        segment_prices: dict[str, object] = {}

        if "Farms" in wb.sheetnames:
            ws = wb["Farms"]
            rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
            idx = _find_header_row(rows, FARMS_HEADER)
            if idx is None:
                issues.append(LoaderIssue(sheet="Farms", message="Could not find the expected header row"))
            else:
                farms = _read_rows_after_header(rows, idx, FARMS_HEADER)

        if "Clients" in wb.sheetnames:
            ws = wb["Clients"]
            rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
            idx = _find_header_row(rows, CLIENTS_HEADER)
            if idx is None:
                issues.append(LoaderIssue(sheet="Clients", message="Could not find the expected header row"))
            else:
                clients = _read_rows_after_header(rows, idx, CLIENTS_HEADER)

        if "Station" in wb.sheetnames:
            ws = wb["Station"]
            rows = [tuple(r) for r in ws.iter_rows(values_only=True)]

            idx = _find_header_row(rows, STATION_HEADER)
            if idx is None:
                issues.append(LoaderIssue(sheet="Station", message="Could not find the station header row"))
            else:
                station_rows = _read_rows_after_header(rows, idx, STATION_HEADER)
                station = station_rows[0] if station_rows else None
                if station is None:
                    issues.append(LoaderIssue(sheet="Station", message="Station header found but no data row follows it"))

            price_idx = _find_header_row(rows, SEGMENT_PRICE_HEADER)
            if price_idx is None:
                issues.append(
                    LoaderIssue(sheet="Station", message="Could not find the segment reference-price table header")
                )
            else:
                for row in _read_rows_after_header(rows, price_idx, SEGMENT_PRICE_HEADER):
                    seg = row.get("segment")
                    if seg is not None:
                        segment_prices[str(seg)] = row.get("reference_export_price_per_t_eur")

        return RawWorkbook(
            farms=farms,
            clients=clients,
            station=station,
            segment_prices=segment_prices,
            loader_issues=issues,
        )
    finally:
        wb.close()
